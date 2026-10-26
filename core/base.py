import io, os, tempfile, subprocess

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx2pdf import convert
from core.subjects import BUILDERS, REPORT_TEMPLATES


TEMPLATES_DIR = os.path.join(os.path.dirname(__file__), '..', 'templates', 'docs')


class ReportGenerator:

    def __init__(self, student=None, tests=None, teacher=None, marks=None):
        self.teacher = teacher
        self.student = student
        self.marks = marks
        self.tests = tests

    # ── Helpers ───────────────────────────────────────────

    def _mark(self, test_type=None, result_pos=0):
        if self.teacher:
            try:
                m = self.marks.filter(student=self.student, test__type=test_type).order_by('weighted', '-test__date')[result_pos]
                return m.mark or ''
            except IndexError:
                return ''

        elif self.student:
            try:
                t = self.tests.filter(type=test_type).order_by('weighted')[result_pos]
                return t.mark or ''
            except IndexError:
                return ''

    def _total(self, test_type=None, result_pos=0):
        if self.teacher:
            try:
                t = self.tests.filter(type=test_type).order_by('-date')[result_pos]
                return t.total or ''
            except IndexError:
                return ''

        elif self.student:
            try:
                t = self.tests.filter(type=test_type).order_by('weighted')[result_pos]
                return t.total or ''
            except IndexError:
                return ''

    def _weighted(self, test_type=None, result_pos=0):
        if self.teacher:
            try:
                t = self.marks.filter(student=self.student, test__type=test_type).order_by('weighted', '-test__date')[result_pos]
                return t.weighted or ''
            except IndexError:
                return ''

        elif self.student:
            try:
                t = self.tests.filter(type=test_type).order_by('weighted')[result_pos]
                return t.weighted or ''
            except IndexError:
                return ''

    def _perc(self, test_type=None, result_pos=0):
        if self.teacher:
            try:
                t = self.marks.filter(student=self.student, test__type=test_type).order_by('weighted', '-test__date')[result_pos]
                return t.percentage or ''
            except IndexError:
                return ''

        elif self.student:
            try:
                t = self.tests.filter(type=test_type).order_by('weighted')[result_pos]
                return t.percentage or ''
            except IndexError:
                return ''

    def _date(self, test_type=None, result_pos=0):
        if self.teacher:
            try:
                t = self.tests.filter(type=test_type).order_by('-date')[result_pos]
                return t.date if t.date is not None else ''
            except IndexError:
                return ''

        elif self.student:
            try:
                t = self.tests.filter(type=test_type).order_by('weighted')[result_pos]
                return str(t.date) if t.date is not None else ''
            except IndexError:
                return ''

    def _desc(self, test_type=None, result_pos=0):
        if self.teacher:
            try:
                t = self.tests.filter(type=test_type).order_by('-date')[result_pos]
                return t.description if t.description is not None else ''
            except IndexError:
                return ''

        elif self.student:
            try:
                t = self.tests.filter(type=test_type).order_by('weighted')[result_pos]
                return t.description if t.description is not None else ''
            except IndexError:
                return ''

    # ── Template engine ───────────────────────────────────

    def _replace_in_paragraph(self, paragraph, data):
        full_text = ''.join(run.text for run in paragraph.runs)

        if not any(key in full_text for key in data):
            return

        for key, value in data.items():
            if key in full_text:
                full_text = full_text.replace(key, str(value))

        paragraph.runs[0].text = full_text
        for run in paragraph.runs[1:]:
            run.text = ''

    def _fill_template(self, template_path, data):
        doc = Document(template_path)

        for paragraph in doc.paragraphs:
            self._replace_in_paragraph(paragraph, data)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    for paragraph in cell.paragraphs:
                        self._replace_in_paragraph(paragraph, data)

        return doc

    def _to_bytes(self, doc):
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer

    # ── Public API ────────────────────────────────────────

    def has_template(self, subject_id):
        return subject_id in REPORT_TEMPLATES

    def generate(self, subject_id):
        template_path = os.path.join(TEMPLATES_DIR, REPORT_TEMPLATES[subject_id])
        data = BUILDERS[subject_id](self)
        doc = self._fill_template(template_path, data)
        return self._to_bytes(doc)

    def generate_pdf(self, subject_id):
        template_path = os.path.join(TEMPLATES_DIR, REPORT_TEMPLATES[subject_id])
        data = BUILDERS[subject_id](self)
        doc = self._fill_template(template_path, data)

        with tempfile.TemporaryDirectory() as tmp:
            docx_path = os.path.join(tmp, 'report.docx')
            pdf_path = os.path.join(tmp, 'report.pdf')
            doc.save(docx_path)

            convert(docx_path, pdf_path)

            with open(pdf_path, 'rb') as f:
                return io.BytesIO(f.read())


